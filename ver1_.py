
# coding: utf-8

# In[5]:

#   Budget Game
#
#   02/22/2018
#   02/25/2018
#
#
import pandas as pd
import numpy as np
#pip.main(['install','openpyxl'])

year = 1
budget = 0
stockpile = 0
available = 0
demand = 0
activity = 0
shortage = 0
surplus = 0
revert = 0
newstock = 0
cost = 0
efficient = 0
effective = 0
score = 0
points = 0
avgscore = 0
newyear = 0
revtInd = 'N'
stkInd = 'N'
newInd = 'N'
newdemand = 0
budget1 = 0
revert1 = 0
finalavg = 0
finalscr = 0
c = 0
cumavg = 0
cavg = 0
clen = 0
Email = ""
Fname = ""
Lname = ""


print(year)

import pandas as pd
import random
import openpyxl
import webbrowser

def act(demand, budget, available):
    if(budget >= demand):
        activity = demand
    elif((budget < demand) and (demand <= available)):
        activity = demand
    elif (available < demand):
        activity = available
    return activity

def short(demand, budget, available):
    if((budget == demand) or (budget > demand)):
        shortage = 0
    elif((budget < demand) and (demand <= available)):
        shortage = demand-budget
    elif(available < demand):
        shortage = demand-available
    return shortage

def surp(demand, budget, available):
    if((budget == demand) or (budget < demand and demand <= available) or (available < demand)):
        surplus = 0
    elif(budget > demand):
        surplus = budget - demand
    return surplus

def revt(demand, budget, available, revtInd):
    if((budget > demand) and (revtInd == 'Y')):
        revert = surplus
    else:
        revert = 0
    return revert


def nws(budget, demand, available, stkInd, revtInd, shortage, surplus,stockpile):
    if(budget == demand):
        newstock = stockpile
    elif((budget > demand) and (stkInd == 'Y')):
        newstock = stockpile + surplus
    elif((budget > demand) and (revtInd == 'Y')):
        newstock = stockpile
    elif((budget < demand) and (demand <= available)):
        newstock = stockpile-shortage
    elif(available < demand):
        newstock = 0
    #else:
     #   newstock = stockpile
    return newstock

def cst(demand, budget, stkInd, revtInd,available, activity):
    if(budget == demand):
        cost = activity
    elif((budget > demand) and (stkInd == 'Y')):
        cost = activity + surplus
    elif((budget > demand) and (revtInd == 'Y')):
        cost = activity
    elif((budget < demand) and (demand <= available)):
        cost = activity
    elif(available < demand):
        cost = activity
    #else:
      #  cost = activity
    return cost

def effi(activity,cost):
    efficient = (float(activity/cost))
    return efficient

def effe(activity,demand):
    effective = (float(activity/demand))
    return effective

def scr(efficient,effective):
    score = ((efficient+effective)/2)
    return score

def pts(points,score):
    points = points + score
    return points

def avg(points, year):
    avgscore = (points/year)
    return avgscore
    
df = pd.DataFrame(index=range(0,60),columns = ['Year', 'Budget', 'Demand','Surplus', 'Shortage','Available', 'Activity', 'Stockpile', 'Revert', 'Newstock', 'Cost', 'Efficient', 'Effective', 'Score', 'Points', 'Avgscore', 'CumAvg', 'EmailID', 'Last_Name', 'First_Name'])

while year < 61:
    if (year == 1 and demand == 0):
        budget = 100
        df['Budget'].iloc[0]= budget
        print("budget is", budget)
        stockpile = 0
        df['Stockpile'].iloc[0]= stockpile
        df['Year'].iloc[0]= year
    
        mylist = [90,100,110]
        demand = random.choice(mylist)
        print("demand is", demand)
        df['Demand'].iloc[0]= demand
        
        available = budget + stockpile
        print("Available is", available)
        df['Available'].iloc[0]= available
        #df[df['Available'].append(available)]
        
        activity = act(demand, budget, available)
        print("activities performed", activity)
        df['Activity'].iloc[0]= activity
        
        #print(df)
        
        shortage = short(demand, budget, available)
        print("shortage is", shortage)
        df['Shortage'].iloc[0]= shortage
        
        surplus = surp(demand, budget, available)
        print("surplus is", surplus)
        df['Surplus'].iloc[0]= surplus
        if(surplus > 0):
            newInd = str(input("Do you want to stockpile the surplus Y/N: "))
            while(newInd not in ("n", "y", "N", "Y")):
                newInd = str(input("Please enter Y/N only, Do you want to stockpile the surplus Y/N: "))
            print(newInd)
            if(newInd == 'Y' or newInd == 'y' ):
                stkInd = 'Y'
                df['Stockpile'].iloc[0]= surplus
            elif(newInd == 'N' or newInd == 'n'):
                revtInd = 'Y'
                print("first surplus reverted in year", year)
            else:
                revtInd = 'N'
        
        if(revtInd == 'Y'):
            revert = revt(demand, budget, available, revtInd)
            print(" first revert is", revert)
            df['Revert'].iloc[0]= revert
        else:
            df['Revert'].iloc[0] = 0
            
        print("Stockpilr indicator is",stkInd)
        newstock = nws(budget, demand, available, stkInd, revtInd, shortage, surplus,stockpile)
        print("newstock is", newstock)
        df['Newstock'].iloc[0]= newstock
        
        cost = cst(demand, budget, stkInd, revtInd,available, activity)
        print("cost is", cost)
        df['Cost'].iloc[0]= cost
        
        efficient = effi(activity,cost)
        print("your efficency is", efficient)
        df['Efficient'].iloc[0]= efficient
        
        effective = effe(activity,demand)
        print("your effectivness", effective)
        df['Effective'].iloc[0]= effective
    
        score=scr(efficient,effective)
        print("your score is", score)
        df['Score'].iloc[0]= score
    
        points = pts(points,score)
        print("your points is", points)
        df['Points'].iloc[0]= points
    
        avgscore = avg(points, year)
        print("your average score is", avgscore)
        df['Avgscore'].iloc[0]= avgscore
        
        cumavg = avgscore
        df['CumAvg'].iloc[0]= cumavg
    
        year = year + 1
        #df['newyear'].iloc[0]= year
        print("year is",year)
#print(df)
    elif(year > 1):
        df['Year'].iloc[year-1] = year
        print("present year is", year)
        if(df['Revert'].iloc[year-2] > 0):
            budget1 = df['Budget'].iloc[year-2]
            print("budget1 is", budget1)
            revert1 = df['Revert'].iloc[year-2]
            print("revert1 is", revert1)
            budget = int(budget1-(0.5*revert1))
            print("since you reverted in last year present budget is", budget)
            df['Budget'].iloc[year-1]= budget
        #if(budget < 50 and budget > 200):
            #budget = 100
            #print("New budget is", budget)
            #df['Budget'].iloc[year-1]= budget
            #print("budget is", budget)
        else:
            newdemand = df['Demand'].iloc[year-2]
            print("new demand is",newdemand)
            mylist1 = [newdemand+10, newdemand, newdemand-10]
            print("my new list is", mylist1)
            budget = random.choice(mylist1)
            print("budget is", budget)
            df['Budget'].iloc[year-1]= budget
        if((budget<50) or (budget>200)):
            budget = 100
            print("New budget is", budget)
            df['Budget'].iloc[year-1]= budget
        
        #stockpile = stockpile + df['Stockpile'].iloc[year-2]
        stockpile = newstock
        df['Stockpile'].iloc[year-1] = stockpile
        print("stockpile is", stockpile)
        
        available = budget + stockpile
        print("available is", available)
        df['Available'].iloc[year-1]= available
        
        mylist2 = [budget+10, budget, budget-10]
        demand = random.choice(mylist2)
        print("demand is", demand)
        df['Demand'].iloc[year-1]= demand
        
        activity = act(demand, budget, available)
        print("activities performed", activity)
        df['Activity'].iloc[year-1]= activity
        
        shortage = short(demand, budget, available)
        print("shortage is", shortage)
        df['Shortage'].iloc[year-1]= shortage
        
        revtInd = 'N'
        stkInd = 'N'
        
        surplus = surp(demand, budget, available)
        print("surplus is", surplus)
        df['Surplus'].iloc[year-1]= surplus
        if(surplus > 0):
            newInd = input("Do you want to stockpile the surplus Y/N: ")
            while(newInd not in ("n", "y", "N", "Y")):
                newInd = str(input("Please enter Y/N only, Do you want to stockpile the surplus Y/N: "))
            print(newInd)
            if(newInd == 'Y' or newInd == 'y'):
                stkInd = 'Y'
                print("surplus stockpiled in year", year)
            elif(newInd == 'N' or newInd == 'n'):
                revtInd = 'Y'
                print("surplus reverted in year", year)
            else:
                revtInd = 'N'
         
        print("check for revert ind", revtInd)
        
        if(revtInd == 'Y'):
            revert = revt(demand, budget, available, revtInd)
            print("revert is", revert)
            df['Revert'].iloc[year-1]= revert
        else:
            df['Revert'].iloc[year-1]= 0
        
        print("Stockpilr indicator is",stkInd)
        
        newstock = nws(budget, demand, available, stkInd, revtInd, shortage, surplus,stockpile)
        print("newstock is", newstock)
        df['Newstock'].iloc[year-1]= newstock
        
        cost = cst(demand, budget, stkInd, revtInd,available, activity)
        print("cost is", cost)
        df['Cost'].iloc[year-1]= cost
        
        efficient = effi(activity,cost)
        df['Efficient'].iloc[year-1]= efficient
        
        effective = effe(activity,demand)
        df['Effective'].iloc[year-1]= effective
    
        score=scr(efficient,effective)
        print("your score is", score*100,"%")
        df['Score'].iloc[year-1]= score
    
        points = pts(points,score)
        print("your points is", points)
        df['Points'].iloc[year-1]= points
    
        avgscore = avg(points, year)
        df['Avgscore'].iloc[year-1]= avgscore
        
        cavg = df['Avgscore'].sum()
        #print("cum avg is",cavg )
        #print("cum len is",year )
        cumavg = cavg /  year
        df['CumAvg'].iloc[year-1]= cumavg
        year = year + 1
        print(year)
        
print("your efficency for the game is", efficient*100,"%")
print("your effectivness for the game is", effective*100,"%")
print("your average score for the game is", avgscore*100,"%")  
print("your cumulative average score for the game is", cumavg*100,"%")   
c = (df['Surplus'] != 0).sum()
d = (df['Shortage'] != 0).sum()
finalavg = df["Avgscore"].mean()
finalscr = df["Score"].mean()

#print(df)

df2 = pd.DataFrame(df,columns=['Year','Demand','Budget'])
df3 = pd.DataFrame(df,columns=['Year','Efficient','Effective'])
df4 = pd.DataFrame(df,columns=['Year','Avgscore', 'CumAvg'])

writer = pd.ExcelWriter('output.xlsx', engine = 'xlsxwriter')
df.to_excel(writer, sheet_name = 'x1')
df2.to_excel(writer, sheet_name = 'x2')
df3.to_excel(writer, sheet_name = 'x3')
df4.to_excel(writer, sheet_name = 'x4')

workbook = writer.book
worksheet = writer.sheets['x2']

chart = workbook.add_chart({'type': 'line'})
#chart.add_series({...})

# Configure a series with a secondary axis
chart.add_series({
    'name':       '=x2!$C$1',
    'categories': '=x2!$B$2:$B$61',
    'values':     '=x2!$C$2:$C$61',
    'y2_axis': 1,
})

chart.add_series({
    'name':       '=x2!$D$1',
    'categories': '=x2!$B$2:$B$61',
    'values':     '=x2!$D$2:$D$61',
})

# Add a chart title and some axis labels.

chart.set_title({'name': 'Budget/Demand v Year'})
chart.set_legend({'position': 'top'})
chart.set_x_axis({'name': 'Year', })
chart.set_y_axis({'name': 'Budget', 'major_gridlines': {'visible': 0}})
chart.set_y_axis({'min': 30, 'max': 210}) #, 'major_gridlines': {'visible': 0}}
#chart.set_y_axis({'interval_unit': 5})
chart.set_y2_axis({'name': 'Demand', 'major_gridlines': {'visible': 0}})
chart.set_y2_axis({'min': 30, 'max': 210})
#chart.set_y2_axis({'interval_unit': 5})
chart.set_size({'width': 1100, 'height': 375})

# Insert the chart into the worksheet (with an offset).
worksheet.insert_chart('E2', chart) #, {'x_offset': 50, 'y_offset': 20}

#workbook.close()

workbook = writer.book
worksheet = writer.sheets['x3']

chart = workbook.add_chart({'type': 'line'})

# Configure a series with a secondary axis
chart.add_series({
    'name':       '=x3!$D$1',
    'categories': '=x3!$B$2:$B$61',
    'values':     '=x3!$D$2:$D$61',
    'y2_axis': 1,
})

chart.add_series({
    'name':       '=x3!$C$1',
    'categories': '=x3!$B$2:$B$61',
    'values':     '=x3!$C$2:$C$61',
})

chart.set_legend({'position': 'top'})

# Add a chart title and some axis labels.
chart.set_title({'name': 'Efficency/Effectivness v Year'})
chart.set_x_axis({'name': 'Year', })
chart.set_y_axis({'name': 'Efficient', 'major_gridlines': {'visible': 0}})
chart.set_y_axis({'min': 0.85, 'max': 1.1}) #, 'major_gridlines': {'visible': 0}}
#chart.set_y_axis({'interval_unit': 0.025})
chart.set_y2_axis({'name': 'Effective', 'major_gridlines': {'visible': 0}})
chart.set_y2_axis({'min': 0.85, 'max': 1.1})
#chart.set_y2_axis({'interval_unit': 0.025})
chart.set_size({'width': 1100, 'height': 350})

# Insert the chart into the worksheet (with an offset).
worksheet.insert_chart('E2', chart, {'x_offset': 50, 'y_offset': 20}) #, {'x_offset': 50, 'y_offset': 20}

workbook = writer.book
worksheet = writer.sheets['x4']

chart = workbook.add_chart({'type': 'line'})

# Configure a series with a secondary axis
chart = workbook.add_chart({'type': 'line'})

# Configure a series with a secondary axis
chart.add_series({
    'name':       '=x4!$D$1',
    'categories': '=x4!$B$2:$B$61',
    'values':     '=x4!$D$2:$D$61',
    'y2_axis': 1,
})

chart.add_series({
    'name':       '=x4!$C$1',
    'categories': '=x4!$B$2:$B$61',
    'values':     '=x4!$C$2:$C$61',
})

chart.set_legend({'position': 'top'})

# Add a chart title and some axis labels.
chart.set_title({'name': 'Avg Score/Cumulative Avg v Year'})
chart.set_x_axis({'name': 'Year', })
chart.set_y_axis({'name': 'CumAvg', 'major_gridlines': {'visible': 0}})
chart.set_y_axis({'min': 0.9, 'max': 1.1}) #, 'major_gridlines': {'visible': 0}}
#chart.set_y_axis({'interval_unit': 0.025})
chart.set_y2_axis({'name': 'AvgScore', 'major_gridlines': {'visible': 0}})
chart.set_y2_axis({'min': 0.9, 'max': 1.1})
#chart.set_y2_axis({'interval_unit': 0.025})
chart.set_size({'width': 1100, 'height': 350})

# Insert the chart into the worksheet (with an offset).
worksheet.insert_chart('E2', chart, {'x_offset': 50, 'y_offset': 20})

workbook.close()

writer.save()


#with pd.ExcelWriter('output.xlsx', engine='xlsxwriter') as writer:    
#df.to_excel(writer, '1st_sheet')   
#df2.to_excel(writer, '2nd_sheet')   
#writer.save()    


#writer = pd.ExcelWriter('output.xlsx',engine='xlsxwriter')
#df.to_excel(writer,sheet_name='Sheet1')
#df2.to_excel(writer,sheet_name='Sheet2')
#writer.save()

#workbook = writer.book
#worksheet = writer.sheets[2nd_sheet]

# Create a new chart object. In this case an embedded chart.

print("your output file has been saved to same location as this executable file")
'''
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
 
fromaddr = "## from address"
toaddr = "## to address"
 
msg = MIMEMultipart()
 
msg['From'] = fromaddr
msg['To'] = toaddr
msg['Subject'] = "Budget Game Result"
 
body = "Hi, This is an Automated mail. Please find the output file attached"
 
msg.attach(MIMEText(body, 'plain'))
 
filename = "output.xlsx"
attachment = open("output.xlsx", "rb")
 
part = MIMEBase('application', 'octet-stream')
part.set_payload((attachment).read())
encoders.encode_base64(part)
part.add_header('Content-Disposition', "attachment; filename= %s" % filename)
 
msg.attach(part)
 
server = smtplib.SMTP('smtp.gmail.com', 587)
server.starttls()
server.login(fromaddr, "## password")
text = msg.as_string()
server.sendmail(fromaddr, toaddr, text)
server.quit()
'''
print("The number of times Surplus occurred in this game", c)
print("The number of times Shortage occured in this game", d)
print("Your Average Score for this game is",finalavg*100,"%")
print("Your final combined efficency and effectiveness for this game is", finalscr*100,"%")

webbrowser.open('output.xlsx')


# In[ ]:



